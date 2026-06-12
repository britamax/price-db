"""Excel I/O service: parse headers, auto-detect columns, stream rows."""
from __future__ import annotations

from pathlib import Path
from typing import Optional

from openpyxl import load_workbook


# Auto-detect column hints (case-insensitive substring match)
NAME_HINTS = ["uraian", "nama", "material", "item", "pekerjaan", "deskripsi", "barang"]
UNIT_HINTS = ["satuan", "unit", "sat", "uom"]
QTY_HINTS = ["volume", "qty", "jumlah", "kuantitas", "quantity", "vol"]


def detect_columns(headers: list[str]) -> dict:
    """Return {name_column, unit_column, qty_column} guesses from header row."""
    def find(hints: list[str]) -> Optional[str]:
        for h in headers:
            hl = (h or "").strip().lower()
            if not hl:
                continue
            for hint in hints:
                if hint in hl:
                    return h
        return None

    return {
        "name_column": find(NAME_HINTS),
        "unit_column": find(UNIT_HINTS),
        "qty_column": find(QTY_HINTS),
    }


def parse_preview(path: Path, max_preview: int = 10) -> dict:
    """Load .xlsx, return {headers, preview_rows[:N], total_rows, detected}."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return {"headers": [], "preview_rows": [], "total_rows": 0, "detected": {}}

    headers = [str(h) if h is not None else "" for h in header_row]

    preview = []
    total = 0
    for i, row in enumerate(rows_iter, 1):
        total += 1
        if i <= max_preview:
            # Convert all values to str/None for safe rendering
            preview.append([("" if v is None else str(v)) for v in row])

    detected = detect_columns(headers)
    return {
        "headers": headers,
        "preview_rows": preview,
        "total_rows": total,
        "detected": detected,
    }
