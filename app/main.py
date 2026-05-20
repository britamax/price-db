import hashlib
import hmac
import os
import secrets
from io import BytesIO, StringIO
from typing import Dict, List, Optional

import csv
from openpyxl import Workbook, load_workbook
from fastapi import Cookie, Depends, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, PlainTextResponse, RedirectResponse, StreamingResponse
from fastapi.encoders import jsonable_encoder

from db import (
    add_category_rule,
    connect_from_env,
    count_raw_rows,
    create_batch,
    delete_category_rule,
    finish_batch,
    get_batch_detail,
    import_rows,
    init_schema,
    list_batches,
    list_category_rules,
    list_raw_rows,
    search_hybrid,
    search_prices,
    stats_summary,
)
from price_utils import detect_columns, parse_supplier_text_prices
from embeddings import embed_single

APP_ADMIN_USER = os.getenv("APP_ADMIN_USER", "admin")
APP_ADMIN_PASSWORD = os.getenv("APP_ADMIN_PASSWORD", "admin")
SESSION_SECRET = os.getenv("SESSION_SECRET", "change-me-please")
SESSION_COOKIE = "pricedb_admin"
ENV = os.environ

app = FastAPI(title="Database Harga Material")


def get_conn():
    return connect_from_env(ENV)


def smart_search(conn, query: str, limit: int = 10, mode: str = "hybrid"):
    """Hybrid search with automatic fallback to lexical if embedding fails."""
    if mode == "lexical":
        results = search_prices(conn, query, limit)
    else:
        try:
            emb = embed_single(query)
            results = search_hybrid(conn, query, emb, limit)
            # Add 'match_score' alias for backward compat with UI/export
            for r in results:
                r["match_score"] = r.get("score")
        except Exception:
            # Fallback to legacy trigram search
            results = search_prices(conn, query, limit)
    # Strip large fields from results
    for r in results:
        r.pop("embedding", None)
        r.pop("embedding_model", None)
        r.pop("embedding_version", None)
        r.pop("embedded_at", None)
    return results


def make_session_token() -> str:
    payload = secrets.token_hex(16)
    sig = hmac.new(SESSION_SECRET.encode(), payload.encode(), hashlib.sha256).hexdigest()
    return f"{payload}.{sig}"


def verify_session_token(token: Optional[str]) -> bool:
    if not token or "." not in token:
        return False
    payload, sig = token.split(".", 1)
    expected = hmac.new(SESSION_SECRET.encode(), payload.encode(), hashlib.sha256).hexdigest()
    return hmac.compare_digest(sig, expected)


def require_admin_cookie(pricedb_admin: Optional[str] = Cookie(default=None)):
    if not verify_session_token(pricedb_admin):
        raise HTTPException(status_code=401, detail="Login admin diperlukan. Buka /admin/login")
    return True


def require_admin(username: str, password: str):
    if username != APP_ADMIN_USER or password != APP_ADMIN_PASSWORD:
        raise HTTPException(status_code=401, detail="Username/password salah")


def html_escape(value) -> str:
    return (
        str(value or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def fmt_rupiah(value) -> str:
    if value is None or value == "":
        return ""
    try:
        return f"Rp {float(value):,.0f}".replace(",", ".")
    except Exception:
        return str(value)


def page(title: str, body: str, admin: bool = False) -> str:
    nav_admin = (
        '<a href="/admin/batches">Batches</a>'
        '<a href="/admin/raw">Raw Rows</a>'
        '<a href="/admin/rules">Kategori</a>'
        '<a href="/admin/upload">Upload</a>'
        '<a href="/admin/text-import">Import Teks</a>'
        '<a href="/admin/logout">Logout</a>'
    ) if admin else '<a href="/admin/login">Admin Login</a>'
    return f"""
    <!doctype html><html lang="id"><head>
    <meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
    <title>{title}</title>
    <style>
      body {{ font-family: Arial, sans-serif; background:#0f172a; color:#e2e8f0; margin:0; padding:24px; }}
      .wrap {{ max-width:1240px; margin:auto; }}
      .card {{ background:#111827; border:1px solid #334155; border-radius:14px; padding:20px; margin:0 0 16px; box-shadow:0 10px 30px #0005; }}
      h1,h2 {{ color:#38bdf8; margin-top:0; }}
      a {{ color:#7dd3fc; }}
      input, textarea, select {{ width:100%; box-sizing:border-box; padding:10px; margin:6px 0 12px; border-radius:8px; border:1px solid #475569; background:#020617; color:#e2e8f0; }}
      button {{ background:#0ea5e9; color:white; border:0; padding:10px 16px; border-radius:8px; cursor:pointer; font-weight:bold; }}
      button.warn {{ background:#f97316; }} button.danger {{ background:#dc2626; }}
      table {{ width:100%; border-collapse:collapse; font-size:13px; }}
      th,td {{ border-bottom:1px solid #334155; padding:8px; vertical-align:top; }}
      th {{ text-align:left; color:#93c5fd; position:sticky; top:0; background:#111827; }}
      .muted {{ color:#94a3b8; }} .ok {{ color:#86efac; }} .warn {{ color:#fde68a; }} .err {{ color:#fca5a5; }}
      .nav {{ display:flex; gap:10px; flex-wrap:wrap; margin-bottom:16px; }}
      .nav a {{ background:#1e293b; padding:8px 12px; border-radius:8px; text-decoration:none; }}
      .stats {{ display:flex; gap:14px; flex-wrap:wrap; }}
      .stat {{ background:#0b1220; border:1px solid #1e293b; border-radius:10px; padding:12px 16px; min-width:150px; }}
      .stat b {{ color:#38bdf8; font-size:18px; }}
      code {{ background:#020617; padding:2px 6px; border-radius:6px; }}
      .pill {{ display:inline-block; padding:2px 8px; border-radius:999px; font-size:11px; }}
      .pill.ok {{ background:#064e3b; color:#bbf7d0; }} .pill.warn {{ background:#713f12; color:#fde68a; }} .pill.err {{ background:#7f1d1d; color:#fecaca; }}
      .row {{ display:grid; grid-template-columns:1fr 1fr; gap:14px; }}
      @media (max-width:720px) {{ .row {{ grid-template-columns:1fr; }} }}
    </style></head><body><div class="wrap">
    <div class="nav"><a href="/">Search</a>{nav_admin}<a href="/health">Health</a></div>
    {body}
    </div>
    <script>
    function initSortable() {{
      document.querySelectorAll('table[class*="sortable"], .sortable-table').forEach(function(tbl) {{
        tbl.querySelectorAll('thead th[data-sort]').forEach(function(th, colIdx) {{
          if (th.dataset.sortBound) return;
          th.dataset.sortBound = '1';
          th.style.cursor = 'pointer';
          th.title = 'Klik untuk sort';
          if (!th.querySelector('.sort-icon')) {{
            var icon = document.createElement('span');
            icon.className = 'sort-icon';
            icon.textContent = ' ⇅';
            icon.style.fontSize = '10px';
            icon.style.opacity = '0.5';
            th.appendChild(icon);
          }}
          th.addEventListener('click', function() {{
            var asc = th.dataset.sortDir !== 'asc';
            th.dataset.sortDir = asc ? 'asc' : 'desc';
            tbl.querySelectorAll('thead th .sort-icon').forEach(function(ic) {{
              ic.textContent = ' ⇅'; ic.style.opacity = '0.5';
            }});
            th.querySelector('.sort-icon').textContent = asc ? ' ▲' : ' ▼';
            th.querySelector('.sort-icon').style.opacity = '1';
            var tbody = tbl.querySelector('tbody');
            var rows = Array.from(tbody.querySelectorAll('tr'));
            rows.sort(function(a, b) {{
              var ca = a.cells[colIdx], cb = b.cells[colIdx];
              var va = (ca.dataset.value !== undefined ? ca.dataset.value : ca.textContent).trim();
              var vb = (cb.dataset.value !== undefined ? cb.dataset.value : cb.textContent).trim();
              var na = parseFloat(va), nb = parseFloat(vb);
              var cmp = (!isNaN(na) && !isNaN(nb)) ? na - nb : va.localeCompare(vb, 'id');
              return asc ? cmp : -cmp;
            }});
            rows.forEach(function(r) {{ tbody.appendChild(r); }});
          }});
        }});
      }});
    }}
    document.addEventListener('DOMContentLoaded', initSortable);
    document.addEventListener('htmx:afterSwap', initSortable);
    </script>
    </body></html>
    """


@app.on_event("startup")
def startup():
    with get_conn() as conn:
        init_schema(conn)
    # Register Phase 3 admin v2 routes
    try:
        from admin_v2 import register_routes
        from embeddings import embed_single as _embed_single
        register_routes(app, get_conn, require_admin_cookie, search_hybrid, _embed_single)
    except Exception as e:
        import traceback
        print(f"[admin_v2] register failed: {e}\n{traceback.format_exc()}")


@app.get("/health", response_class=PlainTextResponse)
def health():
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT COUNT(*) AS c FROM price_records")
            c = cur.fetchone()["c"]
    return f"OK - database connected - {c} price records"


def render_results(results_by_query: Dict[str, List[dict]]) -> str:
    chunks = []
    for query, rows in results_by_query.items():
        chunks.append(f"<h2>Query: <code>{html_escape(query)}</code></h2>")
        if not rows:
            chunks.append("<p class='warn'>Tidak ada hasil.</p>")
            continue
        chunks.append("<div style='overflow:auto'><table class='sortable-table'><thead><tr><th data-sort='score'>Score</th><th data-sort='tanggal'>Tanggal</th><th data-sort='nama'>Nama</th><th data-sort='merek'>Merek</th><th data-sort='spec'>Spec</th><th data-sort='satuan'>Satuan</th><th data-sort='harga'>Harga</th><th data-sort='supplier'>Supplier</th><th data-sort='kategori'>Kategori</th><th data-sort='keterangan'>Keterangan</th></tr></thead><tbody>")
        for r in rows:
            score = r.get("match_score")
            score_txt = f"{float(score):.2f}" if score is not None else ""
            score_val = f"{float(score):.4f}" if score is not None else "0"
            harga_val = str(r.get('harga') or 0)
            chunks.append(
                f"<tr>"
                f"<td data-value='{score_val}'>{score_txt}</td>"
                f"<td>{html_escape(r.get('effective_date'))}</td>"
                f"<td>{html_escape(r.get('nama'))}</td>"
                f"<td>{html_escape(r.get('merek'))}</td>"
                f"<td>{html_escape(r.get('spesifikasi'))}</td>"
                f"<td>{html_escape(r.get('satuan'))}</td>"
                f"<td data-value='{harga_val}'>{fmt_rupiah(r.get('harga'))}</td>"
                f"<td>{html_escape(r.get('supplier'))}</td>"
                f"<td>{html_escape(r.get('category'))}/{html_escape(r.get('subcategory'))}</td>"
                f"<td>{html_escape(r.get('keterangan'))}</td>"
                f"</tr>"
            )
        chunks.append("</tbody></table></div>")
    return "".join(chunks)


def admin_session(pricedb_admin: Optional[str] = Cookie(default=None)) -> bool:
    return verify_session_token(pricedb_admin)


@app.get("/", response_class=HTMLResponse)
def home(pricedb_admin: Optional[str] = Cookie(default=None)):
    body = """
    <div class="card"><h1>Database Harga Material</h1>
    <p class="muted">Cari harga material historis. Bisa satu item atau banyak item, satu baris per query. Hasil terbaru di atas.</p>
    <form action="/search" method="post">
      <label>Query / daftar item</label>
      <textarea name="queries" rows="8" placeholder="kabel nyy 3x4 supreme&#10;semen gresik&#10;pipa pvc 3 inch"></textarea>
      <label>Top hasil per query</label><input name="limit" value="10">
      <button type="submit">Cari</button>
    </form></div>
    """
    return HTMLResponse(page("Search Harga Material", body, admin=admin_session(pricedb_admin)))


@app.post("/search", response_class=HTMLResponse)
def search(queries: str = Form(...), limit: int = Form(10), pricedb_admin: Optional[str] = Cookie(default=None)):
    lines = [l.strip() for l in queries.splitlines() if l.strip()]
    results = {}
    with get_conn() as conn:
        for q in lines:
            results[q] = smart_search(conn, q, max(1, min(limit, 50)))
    body = f"""
    <div class="card"><h1>Hasil Search</h1>
    <form action="/export-search.csv" method="post" style="display:inline-block"><input type="hidden" name="queries" value="{html_escape(queries)}"><input type="hidden" name="limit" value="{limit}"><button>Export CSV</button></form>
    <form action="/export-search.xlsx" method="post" style="display:inline-block;margin-left:8px"><input type="hidden" name="queries" value="{html_escape(queries)}"><input type="hidden" name="limit" value="{limit}"><button>Export Excel</button></form>
    </div><div class="card">{render_results(results)}</div>
    """
    return HTMLResponse(page("Hasil Search", body, admin=admin_session(pricedb_admin)))


@app.get("/api/search")
def api_search(q: str, limit: int = 5, mode: str = "hybrid"):
    safe_limit = max(1, min(limit, 10))
    with get_conn() as conn:
        rows = smart_search(conn, q, safe_limit, mode=mode)
    return {"query": q, "limit": safe_limit, "results": jsonable_encoder(rows)}


@app.post("/api/batch-search")
def api_batch_search(payload: dict):
    queries = payload.get("queries") or []
    if isinstance(queries, str):
        queries = [l.strip() for l in queries.splitlines() if l.strip()]
    queries = [str(q).strip() for q in queries if str(q).strip()][:10]
    safe_limit = max(1, min(int(payload.get("limit", 5)), 10))
    results = {}
    with get_conn() as conn:
        for q in queries:
            results[q] = jsonable_encoder(smart_search(conn, q, safe_limit))
    return {"limit": safe_limit, "results": results}


def _collect_search_rows(queries: str, limit: int):
    cols = ["query","match_score","tanggal","nama","merek","spesifikasi","satuan","harga","supplier","category","subcategory","lokasi","sumber","keterangan"]
    rows = []
    with get_conn() as conn:
        for q in [l.strip() for l in queries.splitlines() if l.strip()]:
            for r in smart_search(conn, q, max(1, min(limit, 50))):
                rows.append([
                    q, r.get("match_score"), r.get("effective_date"), r.get("nama"), r.get("merek"),
                    r.get("spesifikasi"), r.get("satuan"), r.get("harga"), r.get("supplier"),
                    r.get("category"), r.get("subcategory"), r.get("lokasi"), r.get("sumber"), r.get("keterangan"),
                ])
    return cols, rows


@app.post("/export-search.csv")
def export_search_csv(queries: str = Form(...), limit: int = Form(10)):
    cols, rows = _collect_search_rows(queries, limit)
    out = StringIO()
    w = csv.writer(out)
    w.writerow(cols)
    for r in rows:
        w.writerow(["" if v is None else v for v in r])
    return StreamingResponse(iter([out.getvalue()]), media_type="text/csv", headers={"Content-Disposition":"attachment; filename=hasil-search.csv"})


@app.post("/export-search.xlsx")
def export_search_xlsx(queries: str = Form(...), limit: int = Form(10)):
    cols, rows = _collect_search_rows(queries, limit)
    wb = Workbook()
    ws = wb.active
    ws.title = "Hasil Search"
    ws.append(cols)
    for r in rows:
        ws.append([("" if v is None else v) for v in r])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition":"attachment; filename=hasil-search.xlsx"})


@app.get("/admin/login", response_class=HTMLResponse)
def admin_login_form():
    body = """
    <div class="card"><h1>Admin Login</h1>
    <form action="/admin/login" method="post">
      <label>Username</label><input name="username" value="admin" required>
      <label>Password</label><input name="password" type="password" required>
      <button>Login</button>
    </form></div>
    """
    return HTMLResponse(page("Admin Login", body))


@app.post("/admin/login")
def admin_login(username: str = Form(...), password: str = Form(...)):
    require_admin(username, password)
    token = make_session_token()
    resp = RedirectResponse(url="/admin/batches", status_code=303)
    resp.set_cookie(SESSION_COOKIE, token, httponly=True, samesite="lax", max_age=8 * 3600)
    return resp


@app.get("/admin/logout")
def admin_logout():
    resp = RedirectResponse(url="/", status_code=303)
    resp.delete_cookie(SESSION_COOKIE)
    return resp


@app.get("/admin/batches", response_class=HTMLResponse)
def admin_batches(_: bool = Depends(require_admin_cookie)):
    with get_conn() as conn:
        stats = stats_summary(conn)
        batches = list_batches(conn, limit=200)
    rows_html = []
    for b in batches:
        status_pill = f"<span class='pill ok'>{html_escape(b['status'])}</span>"
        rows_html.append(
            f"<tr><td>{b['id']}</td><td>{html_escape(b['uploaded_at'])}</td><td>{html_escape(b['filename'])}</td><td>{html_escape(b['input_source'])}</td><td>{html_escape(b['uploaded_by'])}</td><td>{status_pill}</td><td>{b['total_rows']}</td><td class='ok'>{b['inserted_rows']}</td><td class='warn'>{b['duplicate_rows']}</td><td class='err'>{b['error_rows']}</td><td><a href='/admin/batches/{b['id']}'>detail</a></td></tr>"
        )
    body = f"""
    <div class="card"><h1>Riwayat Upload</h1>
    <div class="stats">
      <div class="stat">Records<br><b>{stats['records']}</b></div>
      <div class="stat">Raw rows<br><b>{stats['raws']}</b></div>
      <div class="stat">Batches<br><b>{stats['batches']}</b></div>
      <div class="stat">Suppliers<br><b>{stats['suppliers']}</b></div>
    </div></div>
    <div class="card"><div style="overflow:auto"><table><thead><tr><th>ID</th><th>Tanggal</th><th>File</th><th>Source</th><th>By</th><th>Status</th><th>Total</th><th>OK</th><th>Dup</th><th>Err</th><th></th></tr></thead><tbody>{''.join(rows_html) or '<tr><td colspan=11>Belum ada batch.</td></tr>'}</tbody></table></div></div>
    """
    return HTMLResponse(page("Batches", body, admin=True))


@app.get("/admin/batches/{batch_id}", response_class=HTMLResponse)
def admin_batch_detail(batch_id: int, _: bool = Depends(require_admin_cookie)):
    with get_conn() as conn:
        batch, events = get_batch_detail(conn, batch_id)
    if not batch:
        return HTMLResponse(page("Batch", "<div class='card'><p class='err'>Batch tidak ditemukan.</p></div>", admin=True))
    rows_html = []
    for ev in events:
        cls = "ok" if ev["status"] == "inserted" else ("warn" if ev["status"] == "duplicate" else "err")
        rows_html.append(
            f"<tr><td>{ev['row_number']}</td><td><span class='pill {cls}'>{html_escape(ev['status'])}</span></td><td>{ev['raw_row_id'] or ''}</td><td>{html_escape(ev['message'])}</td><td>{html_escape(ev['created_at'])}</td></tr>"
        )
    body = f"""
    <div class="card"><h1>Batch #{batch['id']}</h1>
    <p>File: <code>{html_escape(batch['filename'])}</code> • Source: <code>{html_escape(batch['input_source'])}</code> • By: {html_escape(batch['uploaded_by'])} • {html_escape(batch['uploaded_at'])}</p>
    <p>Status: <b>{html_escape(batch['status'])}</b> • Total: {batch['total_rows']} • Inserted: {batch['inserted_rows']} • Duplicate: {batch['duplicate_rows']} • Error: {batch['error_rows']}</p>
    <p>Notes: {html_escape(batch.get('notes'))}</p></div>
    <div class="card"><h2>Events</h2><div style="overflow:auto"><table><thead><tr><th>Row</th><th>Status</th><th>Raw ID</th><th>Message</th><th>Time</th></tr></thead><tbody>{''.join(rows_html) or '<tr><td colspan=5>(no events)</td></tr>'}</tbody></table></div></div>
    """
    return HTMLResponse(page(f"Batch {batch_id}", body, admin=True))


@app.get("/admin/raw", response_class=HTMLResponse)
def admin_raw(q: str = "", page_no: int = 1, _: bool = Depends(require_admin_cookie)):
    page_no = max(1, page_no)
    page_size = 50
    offset = (page_no - 1) * page_size
    with get_conn() as conn:
        rows = list_raw_rows(conn, search=q, limit=page_size, offset=offset)
        total = count_raw_rows(conn)
    rows_html = []
    for r in rows:
        rows_html.append(
            f"<tr><td>{r['id']}</td><td>{html_escape(r['nama_raw'])}</td><td>{html_escape(r['merek_raw'])}</td><td>{html_escape(r['satuan_mapped'])}</td><td>{fmt_rupiah(r['harga_parsed'])}</td><td>{html_escape(r['supplier_raw'])}</td><td>{html_escape(r['tanggal_parsed'])}</td><td>{r['seen_count']}</td><td>{r['first_batch_id']}</td><td>{html_escape(r['created_at'])}</td></tr>"
        )
    body = f"""
    <div class="card"><h1>Raw Rows</h1>
    <p class="muted">Total: {total}. Halaman {page_no}.</p>
    <form action="/admin/raw" method="get"><input name="q" value="{html_escape(q)}" placeholder="cari nama/merek/supplier"><button>Cari</button></form>
    </div>
    <div class="card"><div style="overflow:auto"><table><thead><tr><th>ID</th><th>Nama</th><th>Merek</th><th>Satuan</th><th>Harga</th><th>Supplier</th><th>Tanggal</th><th>Seen</th><th>1st batch</th><th>Created</th></tr></thead><tbody>{''.join(rows_html) or '<tr><td colspan=10>Belum ada data.</td></tr>'}</tbody></table></div>
    <p><a href="/admin/raw?q={html_escape(q)}&page_no={max(1,page_no-1)}">« prev</a> | <a href="/admin/raw?q={html_escape(q)}&page_no={page_no+1}">next »</a></p></div>
    """
    return HTMLResponse(page("Raw Rows", body, admin=True))


@app.get("/admin/rules", response_class=HTMLResponse)
def admin_rules(_: bool = Depends(require_admin_cookie)):
    with get_conn() as conn:
        rules = list_category_rules(conn)
    rows_html = []
    for r in rules:
        rows_html.append(
            f"<tr><td>{r['id']}</td><td><code>{html_escape(r['keyword'])}</code></td><td>{html_escape(r['category'])}</td><td>{html_escape(r['subcategory'])}</td><td>{r['priority']}</td><td>{'on' if r['active'] else 'off'}</td><td><form action='/admin/rules/{r['id']}/delete' method='post' style='display:inline'><button class='danger'>Hapus</button></form></td></tr>"
        )
    body = f"""
    <div class="card"><h1>Rule Kategori</h1>
    <p class="muted">Keyword cocok jika ditemukan di nama (lowercase, partial match). Priority lebih kecil = lebih dulu dicek.</p>
    <form action="/admin/rules" method="post" class="row">
      <div><label>Keyword</label><input name="keyword" required placeholder="kabel nyy"></div>
      <div><label>Priority</label><input name="priority" value="100"></div>
      <div><label>Category</label><input name="category" required placeholder="elektrikal"></div>
      <div><label>Subcategory</label><input name="subcategory" required placeholder="kabel"></div>
      <div style="grid-column:1/-1"><button>Tambah / Update</button></div>
    </form></div>
    <div class="card"><div style="overflow:auto"><table><thead><tr><th>ID</th><th>Keyword</th><th>Category</th><th>Subcategory</th><th>Priority</th><th>Active</th><th></th></tr></thead><tbody>{''.join(rows_html) or '<tr><td colspan=7>Belum ada rule.</td></tr>'}</tbody></table></div></div>
    """
    return HTMLResponse(page("Rule Kategori", body, admin=True))


@app.post("/admin/rules", response_class=HTMLResponse)
def admin_rules_add(keyword: str = Form(...), category: str = Form(...), subcategory: str = Form(...), priority: int = Form(100), _: bool = Depends(require_admin_cookie)):
    with get_conn() as conn:
        add_category_rule(conn, keyword, category, subcategory, priority)
    return RedirectResponse(url="/admin/rules", status_code=303)


@app.post("/admin/rules/{rule_id}/delete")
def admin_rules_delete(rule_id: int, _: bool = Depends(require_admin_cookie)):
    with get_conn() as conn:
        delete_category_rule(conn, rule_id)
    return RedirectResponse(url="/admin/rules", status_code=303)


@app.get("/admin/upload", response_class=HTMLResponse)
def upload_form(_: bool = Depends(require_admin_cookie)):
    body = """
    <div class="card"><h1>Upload Excel Harga</h1>
    <p class="muted">Kolom minimal: nama, harga. Disarankan: merek, satuan, grup, supplier, tanggal, keterangan.</p>
    <form action="/admin/upload" method="post" enctype="multipart/form-data">
      <label>Catatan batch</label><textarea name="notes"></textarea>
      <label>File Excel/CSV</label><input name="file" type="file" required>
      <button type="submit">Upload</button>
    </form></div>
    """
    return HTMLResponse(page("Upload Excel", body, admin=True))


@app.post("/admin/upload", response_class=HTMLResponse)
async def upload_excel(notes: str = Form(""), file: UploadFile = File(...), _: bool = Depends(require_admin_cookie), pricedb_admin: Optional[str] = Cookie(default=None)):
    content = await file.read()
    file_hash = hashlib.sha256(content).hexdigest()
    records = []
    if file.filename.lower().endswith(".csv"):
        decoded = content.decode("utf-8-sig")
        reader = csv.DictReader(StringIO(decoded))
        columns = reader.fieldnames or []
        records = list(reader)
    else:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        columns = [str(c or "").strip() for c in next(rows_iter)]
        for values in rows_iter:
            records.append({columns[i]: values[i] if i < len(values) else "" for i in range(len(columns))})
    mapping = detect_columns(columns)
    rows = []
    for rec in records:
        row = {target: ("" if rec.get(col) is None else rec.get(col)) for target, col in mapping.items()}
        rows.append(row)
    with get_conn() as conn:
        batch_id = create_batch(conn, file.filename, "admin", "web_upload", notes, file_hash)
        stats = import_rows(conn, batch_id, rows)
        finish_batch(conn, batch_id, "done", stats["total"], stats["inserted"], stats["duplicate"], stats["errors"])
    body = f"""
    <div class="card"><h1>Upload selesai</h1>
    <p>Batch ID: <code>{batch_id}</code></p>
    <p>Kolom terdeteksi: <code>{html_escape(mapping)}</code></p>
    <p class="ok">Inserted: {stats['inserted']}</p><p class="warn">Duplicate: {stats['duplicate']}</p><p class="err">Error: {stats['errors']}</p><p>Total: {stats['total']}</p>
    <p><a href="/admin/batches/{batch_id}">Lihat detail batch</a> | <a href="/">Ke search</a></p></div>
    """
    return HTMLResponse(page("Upload selesai", body, admin=True))


def _render_text_preview(rows, supplier, effective_date, general_note, text):
    rows_html = []
    for i, r in enumerate(rows, 1):
        rows_html.append(
            f"<tr><td>{i}</td><td>{html_escape(r.get('nama'))}</td><td>{html_escape(r.get('merek'))}</td><td>{html_escape(r.get('satuan'))}</td><td>{fmt_rupiah(r.get('harga'))}</td><td>{html_escape(r.get('keterangan'))}</td></tr>"
        )
    return f"""
    <div class="card"><h1>Preview Import Teks</h1>
    <p>Supplier: <b>{html_escape(supplier)}</b> • Tanggal: <b>{html_escape(effective_date)}</b></p>
    <p>Catatan umum: {html_escape(general_note)}</p>
    <p class="muted">{len(rows)} baris hasil parsing. Periksa sebelum commit.</p>
    <div style="overflow:auto"><table><thead><tr><th>#</th><th>Nama</th><th>Merek</th><th>Satuan</th><th>Harga</th><th>Keterangan</th></tr></thead><tbody>{''.join(rows_html) or '<tr><td colspan=6>Tidak ada baris terparsing.</td></tr>'}</tbody></table></div>
    <form action="/admin/text-import/commit" method="post" style="margin-top:14px">
      <input type="hidden" name="supplier" value="{html_escape(supplier)}">
      <input type="hidden" name="effective_date" value="{html_escape(effective_date)}">
      <input type="hidden" name="general_note" value="{html_escape(general_note)}">
      <textarea name="text" hidden>{html_escape(text)}</textarea>
      <button>Commit ke database</button>
      <a href="/admin/text-import" style="margin-left:10px">Batal / edit lagi</a>
    </form></div>
    """


@app.get("/admin/text-import", response_class=HTMLResponse)
def text_import_form(_: bool = Depends(require_admin_cookie)):
    body = """
    <div class="card"><h1>Import Teks Supplier</h1>
    <p class="warn">Klik "Preview" dulu untuk review hasil parsing sebelum masuk database.</p>
    <form action="/admin/text-import/preview" method="post">
      <label>Supplier</label><input name="supplier" placeholder="Sinar Surabaya Sakti" required>
      <label>Tanggal harga</label><input name="effective_date" placeholder="Jumat 8 Mei 2026" required>
      <label>Catatan umum</label><textarea name="general_note">harga include PPN, tidak mengikat</textarea>
      <label>Teks harga supplier</label><textarea name="text" rows="14" placeholder="Kabel NYA 1.5 mm — Rp 413.000/roll @100m (supreme) / Rp 402.000/roll @100m (kabelindo)"></textarea>
      <button type="submit">Preview</button>
    </form></div>
    """
    return HTMLResponse(page("Import Teks Supplier", body, admin=True))


@app.post("/admin/text-import/preview", response_class=HTMLResponse)
def text_import_preview(supplier: str = Form(...), effective_date: str = Form(...), general_note: str = Form(""), text: str = Form(...), _: bool = Depends(require_admin_cookie)):
    rows = parse_supplier_text_prices(text, supplier=supplier, effective_date=effective_date, general_note=general_note)
    return HTMLResponse(page("Preview Import", _render_text_preview(rows, supplier, effective_date, general_note, text), admin=True))


@app.post("/admin/text-import/commit", response_class=HTMLResponse)
def text_import_commit(supplier: str = Form(...), effective_date: str = Form(...), general_note: str = Form(""), text: str = Form(...), _: bool = Depends(require_admin_cookie)):
    rows = parse_supplier_text_prices(text, supplier=supplier, effective_date=effective_date, general_note=general_note)
    with get_conn() as conn:
        batch_id = create_batch(conn, f"text-{supplier}", "admin", "text_import", general_note)
        stats = import_rows(conn, batch_id, rows)
        finish_batch(conn, batch_id, "done", stats["total"], stats["inserted"], stats["duplicate"], stats["errors"])
    body = f"""
    <div class="card"><h1>Import teks selesai</h1>
    <p>Parsed rows: {len(rows)}</p><p>Batch ID: <code>{batch_id}</code></p>
    <p class="ok">Inserted: {stats['inserted']}</p><p class="warn">Duplicate: {stats['duplicate']}</p><p class="err">Error: {stats['errors']}</p>
    <p><a href="/admin/batches/{batch_id}">Lihat detail batch</a> | <a href="/">Ke search</a></p></div>
    """
    return HTMLResponse(page("Import Teks selesai", body, admin=True))
