"""Batch Excel matching worker.

Reads input .xlsx, for each row runs hybrid search, writes output xlsx with
top-1 match + 2 alternatives. Updates batch_job.processed every 10 rows for
HTMX polling.
"""
from __future__ import annotations

import os
import time
import traceback
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook, load_workbook


# Output columns appended after the original sheet's columns
OUTPUT_COLS = [
    "harga_match", "nama_match", "satuan_match", "score", "supplier", "tanggal",
    "alt1_nama", "alt1_harga", "alt1_score",
    "alt2_nama", "alt2_harga", "alt2_score",
]


def _update_progress(conn, job_id: int, processed: int, total: Optional[int] = None):
    with conn.cursor() as cur:
        if total is not None:
            cur.execute(
                "UPDATE batch_job SET processed = %s, total_rows = %s WHERE id = %s",
                (processed, total, job_id),
            )
        else:
            cur.execute("UPDATE batch_job SET processed = %s WHERE id = %s", (processed, job_id))


def _mark_status(conn, job_id: int, status: str, **kwargs):
    sets = ["status = %s"]
    params: list = [status]
    if status == "running":
        sets.append("started_at = now()")
    if status in ("done", "failed", "cancelled"):
        sets.append("finished_at = now()")
    for k, v in kwargs.items():
        sets.append(f"{k} = %s")
        params.append(v)
    params.append(job_id)
    with conn.cursor() as cur:
        cur.execute(f"UPDATE batch_job SET {', '.join(sets)} WHERE id = %s", tuple(params))


def _search_one(conn, query: str, embed_single, search_hybrid):
    """Run hybrid search for one row's name. Returns up to 3 matches."""
    if not query or not query.strip():
        return []
    try:
        emb = embed_single(query)
        results = search_hybrid(conn, query, emb, limit=3)
        return results
    except Exception:
        return []


async def run_batch_match(ctx: dict, job_id: int) -> dict:
    """ARQ entrypoint. Process a batch_job row by row."""
    # Lazy import to avoid circular & ensure runtime context
    import os
    from db import connect_from_env, search_hybrid
    from embeddings import embed_single

    def get_conn():
        return connect_from_env(dict(os.environ))

    started = time.time()
    upload_dir = Path("/app/batch-uploads")

    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT id, filename, original_name, name_column, unit_column, qty_column "
                "FROM batch_job WHERE id = %s",
                (job_id,),
            )
            job = cur.fetchone()
        if not job:
            return {"error": f"job {job_id} not found"}

        input_path = upload_dir / job["filename"]
        if not input_path.exists():
            _mark_status(conn, job_id, "failed", error=f"input file missing: {input_path}")
            return {"error": "input missing"}

        try:
            _mark_status(conn, job_id, "running")
            # Load workbook (read-only mode for memory safety on large files)
            wb_in = load_workbook(input_path, read_only=True, data_only=True)
            ws_in = wb_in.active

            # Materialize rows (small batches expected)
            all_rows = list(ws_in.iter_rows(values_only=True))
            if not all_rows:
                _mark_status(conn, job_id, "failed", error="empty workbook")
                return {"error": "empty"}

            headers = [str(h) if h is not None else "" for h in all_rows[0]]
            data_rows = all_rows[1:]
            total = len(data_rows)
            _update_progress(conn, job_id, 0, total=total)

            # Resolve column indices
            def col_idx(name: Optional[str]) -> Optional[int]:
                if not name:
                    return None
                for i, h in enumerate(headers):
                    if h.strip().lower() == name.strip().lower():
                        return i
                return None

            name_idx = col_idx(job["name_column"])
            unit_idx = col_idx(job["unit_column"])
            qty_idx = col_idx(job["qty_column"])

            if name_idx is None:
                _mark_status(conn, job_id, "failed", error=f"name column '{job['name_column']}' not found")
                return {"error": "name col missing"}

            # Prepare output workbook
            wb_out = Workbook()
            ws_out = wb_out.active
            ws_out.append(headers + OUTPUT_COLS)

            for i, row in enumerate(data_rows, 1):
                row_list = list(row)
                # Pad to header length
                while len(row_list) < len(headers):
                    row_list.append(None)

                query_name = row_list[name_idx]
                query_str = str(query_name).strip() if query_name else ""

                results = _search_one(conn, query_str, embed_single, search_hybrid) if query_str else []

                # Top-1 (confidence threshold to avoid garbage matches)
                MIN_SCORE = 0.30
                if results and (results[0].get("score") or 0) >= MIN_SCORE:
                    r0 = results[0]
                    out_row = row_list + [
                        float(r0["harga"]) if r0.get("harga") is not None else None,
                        r0.get("nama"),
                        r0.get("satuan"),
                        round(float(r0.get("score") or 0), 4),
                        r0.get("supplier"),
                        str(r0.get("effective_date") or "") or None,
                    ]
                else:
                    out_row = row_list + [None, "[no confident match]" if results else None,
                                          None, round(float(results[0].get("score") or 0), 4) if results else None,
                                          None, None]

                # Alt-1 & Alt-2
                for j in (1, 2):
                    if len(results) > j:
                        r = results[j]
                        out_row += [
                            r.get("nama"),
                            float(r["harga"]) if r.get("harga") is not None else None,
                            round(float(r.get("score") or 0), 4),
                        ]
                    else:
                        out_row += [None, None, None]

                ws_out.append(out_row)

                # Update progress every 10 rows (or last row)
                if i % 10 == 0 or i == total:
                    _update_progress(conn, job_id, i)

            # Save output
            out_filename = f"result_{job_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            out_path = upload_dir / out_filename
            wb_out.save(out_path)

            _mark_status(conn, job_id, "done", result_path=str(out_filename))
            return {
                "job_id": job_id,
                "total": total,
                "duration_s": round(time.time() - started, 2),
                "output": out_filename,
            }
        except Exception as e:
            err = f"{type(e).__name__}: {e}\n{traceback.format_exc()[:1500]}"
            _mark_status(conn, job_id, "failed", error=err)
            return {"error": err}
